"""
PR Score Loft — Script de automação
Uso: python pr_score.py planilha_clipping.xlsx
"""

import sys
import json
import re
import time
from datetime import datetime, date, timedelta

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font
import anthropic


# ── PREENCHIMENTO DE CORES ────────────────────────────────────────────────────
RED_FILL    = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="F05A28", end_color="F05A28", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


# ── LISTA TIER 1 ──────────────────────────────────────────────────────────────
TIER1 = {
    "98 FM", "Agência Brasil (EBC)", "Agência Estado", "Agência O Globo",
    "Alô Alô Bahia", "Antagonista", "O Antagonista", "Baguete", "Band",
    "Band News", "Band RS", "Band TV RS", "BandNews", "BandNews FM",
    "BandNews FM - Curitiba", "BandNews FM Curitiba", "BandNews FM (Curitiba)",
    "BBC Brasil", "BC Notícias", "Bloomberg", "Bloomberg Línea", "Brazil Journal",
    "Business Insider", "Business Rock", "Canal Rural", "Canaltech",
    "Capital Aberto", "Casa Claudia", "Casa e Jardim", "Casa Vogue",
    "CBN", "CBN (Nacional)", "CBN Campinas", "CBN Curitiba", "CBN Porto Alegre",
    "CBN Santos", "CNBC - Times Brasil", "CNN", "CNN Brasil",
    "Correio (BA)", "Correio Braziliense", "Correio do Povo",
    "Correio do Povo (RS)", "Correio Popular", "DC (NSC)", "Destak",
    "Diário Catarinense", "Diário da Cidade", "Diário do Comércio",
    "Diário do Grande ABC", "Diário do Nordeste", "Diário do Pará",
    "Economia SC", "EPTV/Globo", "Estado de Minas", "Estadão",
    "Exame", "Exame.com", "Extra", "Fast Company", "Financial Times",
    "Folha de S. Paulo", "Folha de São Paulo", "Folha Dirigida",
    "Forbes", "Forbes Brasil", "Fortune", "G1", "G1 - Campinas",
    "G1 - Grande ABC", "G1 - Interior", "G1 - Nacional", "G1 - Piracicaba",
    "G1 - Porto Alegre", "G1 - Ribeirão e Franca", "G1 - Rio", "G1 - RS",
    "G1 - RS (RBS)", "G1 - Santos", "G1 - São José do Rio Preto",
    "G1 - São Paulo", "G1 - Sorocaba", "G1 - Sul de Minas",
    "G1 - Vale do Paraíba", "G1 - Vanguarda", "G1 (Nacional)",
    "G1 Campinas e Região", "G1 Santa Catarina (NSC)", "G1 Santos e Região",
    "Gaúcha ZH", "GauchazH", "Gazeta do Povo", "Gazeta do Povo (PR)",
    "Gizmodo", "Globo News", "Globonews", "Globo.com", "GQ Brasil",
    "GZH", "GZH (Gaúcha ZH)", "Hoje em Dia", "Hora de Santa Catarina",
    "Imobi Report", "Imobireport", "InfoMoney", "InvestNews",
    "Isto é", "IstoÉ", "Istoé Dinheiro", "Isto é dinheiro",
    "JB FM (Rio de Janeiro)", "Jornal do Comércio",
    "Jornal do Comércio - Porto Alegre", "Jornal Meia Hora", "Jornal Metro",
    "Jornal O Pioneiro", "Jornal O Sul", "Jornal O Sul (online)",
    "Jornal O Tempo", "Jornal O Vale", "Jornal Pequeno",
    "Jota", "JOTA", "Jovem Pan", "Jovem Pan News", "Jovem Pan RS", "JP News",
    "LinkedIn News", "Meio & Mensagem", "Meio e Mensagem", "Meio&Mensagem",
    "Revista Meio & Mensagem - SP", "Meon", "Metro Quadrado", "Metrópoles",
    "MIT Technology Review", "Mobile Time", "Monitor Mercantil", "Money Times",
    "MSN", "MSN Brasil", "Mundo do Marketing", "Na Telinha", "Nd Mais",
    "Neo Feed", "NeoFeed", "Nexo", "NSC Total", "NSC TV", "NSC TV (Globo)",
    "NSC TV (SC)", "NY Times", "O Antagonista", "O Dia", "O Estado de S.Paulo",
    "O Estado de São Paulo", "O Globo", "O Metro SP", "O Sul",
    "O Sul - Porto Alegre", "O Tempo", "Olhar Digital", "Omelete", "OPOVO",
    "Panorama TV", "PEGN", "Pequenas Empresas & Grandes Negócios", "Pioneiro",
    "Portal do Bitcoin", "Portal G1", "Portal Metrópoles",
    "Portal Metrópole Online", "Propmark", "Publicitários Criativos",
    "Quartz", "R7", "Rádio 102.3", "Rádio Band News FM", "Rádio BandNews FM",
    "Rádio Bandeirantes (94.9 FM)", "Rádio Caiçara (Rede Pampa)",
    "Radio Caiçara", "Rádio CBN Porto Alegre", "Rádio Gaúcha", "Rádio Guaíba",
    "Rádio Itatiaia", "Rádio Pampa", "Radio Pampa", "Rádio Super Notícia",
    "Rádio Tupi (96.5 FM - Rio de Janeiro)", "Super Rádio Tupi - 96,5 FM",
    "RBS TV", "Record News", "Record TV", "Record TV RS", "Reuters",
    "SBT", "SBT (Nacional)", "SBT RS", "SCC10", "Santa (Itajaí-NSC)",
    "Santa Revista-SC", "StartSe", "Startups", "Suno", "Suno Notícias",
    "Superinteressante", "TechCrunch", "Tech Crunch", "Tech Tudo", "TechTudo",
    "Tec Mundo", "TecMundo", "Terra", "Times Brasil", "Tribuna do Paraná",
    "TV Band", "TV Cultura", "TV Globo", "TV Pampa", "TV Record RS",
    "TV Tribuna", "Uai", "Uol", "Valor", "Valor Econômico", "Valor Investe",
    "Venture Capital Journal", "Veja", "Veja Rio", "Veja São Paulo",
    "Wall Street Journal", "Washington Post", "Wired", "Yahoo", "YouTube",
    "Zero Hora", "www.radiocaicara.com.br", "Época", "Época Negócios",
}

PAYWALL = {
    "Valor Econômico", "Jornal do Comércio", "Jornal do Comércio - Porto Alegre",
    "Financial Times", "Wall Street Journal", "Washington Post", "NY Times",
}

BRAND_MAP = {
    "grupo loft / loft": "Loft",
    "grupo loft / foxter": "Loft",
    "concorrentes nacionais / zap": "ZAP",
    "concorrentes nacionais / olx": "ZAP",
    "concorrentes nacionais / viva real": "ZAP",
    "concorrentes nacionais / zap imóveis": "ZAP",
    "concorrentes nacionais / quinto andar": "QuintoAndar",
    "concorrentes nacionais / imovelweb": "QuintoAndar",
    "concorrentes nacionais / creditas": "Creditas",
    "concorrentes nacionais / superlógica": "Superlógica",
    "concorrentes nacionais / porto seguro": "Porto Seguro",
    "concorrentes nacionais / lais": "Lais",
    "concorrentes nacionais / lastro": "Lais",
    "concorrentes nacionais / credaluga": "CredAluga",
    "concorrentes nacionais / credpronto": "Credpronto",
    "concorrentes nacionais / credimorar": "Credimorar",
    "concorrentes nacionais / kenlo": "Kenlo",
    "concorrentes nacionais / cashgo": "CashGo",
}

GAB_HEADER = [
    "Índice", "Mês", "Título", "Veículo", "Data", "Link (Clipadora)",
    "Empresa", "Protagonismo", "Critério que pontuou (1-7)",
    "C8 Editorial", "C9 Menções", "C10 CTA", "C11 Visuais",
    "C12 Gráficos", "C13 Páginas", "C14 Capas/Redes",
    "C15 Audiovisual", "C16 Análise Especial",
    "Soma", "PR Tec+Produto", "PR Puro Imobis",
    "Data Proativa", "Data Carona", "Retranca",
    "Tema", "Subtema", "Produto",
    "OBS — Ações necessárias", "Cidade", "Estado"
]

PROMPT_AVALIACAO = """Você é especialista em avaliação de PR Score para o mercado imobiliário brasileiro.

EMPRESA MONITORADA: {empresa}
TÍTULO: {titulo}
VEÍCULO: {veiculo}
TIPO DE MÍDIA: {midia}

TEXTO DA MATÉRIA:
\"\"\"
{texto}
\"\"\"

REGRAS DE AVALIAÇÃO:

PASSO 1 — Verificar se a matéria pontua (critérios 1-7):
Se ALGUM for verdadeiro → protagonismo = "Destaque" (1pt base). Se NENHUM → "Menção" (0pt), encerrar.

1. Porta-voz da empresa com 3+ aspas diretas OU artigo assinado integralmente pelo porta-voz
2. 5+ frases completas OU 8+ linhas referenciando a marca/porta-voz ao longo do texto
3. Marca ocupa 30%+ do volume total da matéria
4. Marca é fonte explícita dos dados/pesquisas da matéria
5. Nota jornalística até 300 toques onde a marca é protagonista
6. Live com 1.000+ visualizações [sinalizar: "ver humano"]
7. Anúncio de governo federal/estadual/prefeitura de capital com empresa citada

PASSO 2 — Pontos extras cumulativos (critérios 8-16, +1 cada):
8. Protagonismo editorial: +1 por local no bloco de abertura [título/chapéu/subtítulo/linha fina — máx +4]
9. 6+ menções nominais à marca no texto
10. CTA explícito "leia mais/veja também" linkando outra matéria da marca (hiperlinks em palavras NÃO contam)
11. Tipos diferentes de foto [legenda/logo/executivo] — "ver humano" se não confirmável pelo texto
12. Tipos diferentes de elemento gráfico [tabela/gráfico/box/olho/lista] — "ver humano" se não confirmável
13. Páginas impressas extras com destaque da marca (só para impresso)
14. Capas e redes sociais → SEMPRE "ver humano"
15. {c15_instrucao}
16. 10+ trechos de porta-voz no lead → SEMPRE "ver humano"

REGRAS OBRIGATÓRIAS:
- pr_tec_produto e pr_puro_imobis SOMENTE "true" para empresa Loft
- soma = 1 (se Destaque) + soma numérica de c8 a c13 (ignorar "ver humano" no cálculo)
- obs: para cada "ver humano", escreva UMA linha direta e acionável. Exemplos:
  "C14: verificar se o veículo postou esta matéria nas redes sociais (Instagram, LinkedIn, X, Facebook, Threads)."
  "C15: matéria de rádio/TV — verificar duração total e número de sonoras do porta-voz."
  "C11: texto menciona foto — confirmar se é foto de executivo, com logo ou com legenda da empresa."
  "C12: texto menciona gráfico/tabela — confirmar se a fonte é a empresa monitorada."
  "C16: matéria com muitos trechos — aplicar protocolo de contagem de sentenças positivas/negativas."
  Se não houver nenhuma sinalização, deixar obs vazio "".

Responda SOMENTE com JSON válido, sem texto adicional:
{{"protagonismo":"Destaque|Menção|Destaque Negativo|Menção Negativa","crit_ativado":"1-7 ou nenhum","c8":0,"c9":0,"c10":0,"c11":0,"c12":0,"c13":0,"c14":"ver humano","c15":{c15_valor},"c16":"ver humano","soma":0,"pr_tec_produto":false,"pr_puro_imobis":false,"data_proativa":false,"data_carona":false,"tema":"","subtema":"","produto":"","obs":"","cidade":"","confianca":"alta|media|baixa"}}"""

HEADERS_HTTP = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.8",
}


def excel_date(v):
    try:
        if isinstance(v, (int, float)):
            d = date(1899, 12, 30) + timedelta(days=int(v))
            return d.strftime("%d/%m/%Y")
        if hasattr(v, "strftime"):
            return v.strftime("%d/%m/%Y")
        return str(v) if v else ""
    except Exception:
        return str(v) if v else ""


def fetch_article(url: str, timeout: int = 15) -> str:
    if not url or not url.startswith("http"):
        return ""
    try:
        resp = requests.get(url, headers=HEADERS_HTTP, timeout=timeout)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        for tag in soup(["script", "style", "nav", "header", "footer",
                          "aside", "form", "iframe", "noscript", "figure"]):
            tag.decompose()
        article = (
            soup.find("article") or
            soup.find(class_=lambda c: c and any(
                k in str(c).lower() for k in ["article", "content", "news", "materia", "texto"]
            )) or
            soup.find("main") or
            soup.body
        )
        text = article.get_text(separator=" ", strip=True) if article else soup.get_text(separator=" ", strip=True)
        text = re.sub(r"\s+", " ", text).strip()
        return text[:5000]
    except Exception:
        return ""


def avaliar_materia(client, row_data, empresa, veiculo, midia, url_fonte, link):
    # Paywall — sinalizar para humano
    if veiculo in PAYWALL:
        return {
            "protagonismo": "Menção",
            "crit_ativado": "nenhum",
            "c8": 0, "c9": 0, "c10": 0, "c11": 0, "c12": 0, "c13": 0,
            "c14": "ver humano",
            "c15": "ver humano" if re.search(r'rádio|radio|tv|podcast', midia, re.I) else 0,
            "c16": "ver humano",
            "soma": 0,
            "pr_tec_produto": False, "pr_puro_imobis": False,
            "data_proativa": False, "data_carona": False,
            "tema": "", "subtema": "", "produto": "",
            "obs": f"PAYWALL: '{veiculo}' tem paywall — acessar via login premium e avaliar todos os critérios manualmente.",
            "cidade": "", "confianca": "baixa",
        }

    texto = fetch_article(url_fonte or link)

    if not texto or len(texto) < 80:
        return {
            "protagonismo": "Menção",
            "crit_ativado": "nenhum",
            "c8": 0, "c9": 0, "c10": 0, "c11": 0, "c12": 0, "c13": 0,
            "c14": "ver humano", "c15": 0, "c16": "ver humano",
            "soma": 0,
            "pr_tec_produto": False, "pr_puro_imobis": False,
            "data_proativa": False, "data_carona": False,
            "tema": "", "subtema": "", "produto": "",
            "obs": f"Não foi possível ler o conteúdo da URL da Fonte. Acessar manualmente: {url_fonte or link}",
            "cidade": "", "confianca": "baixa",
        }

    is_av = bool(re.search(r'rádio|radio|\btv\b|televisão|podcast', midia, re.I))
    c15_instrucao = "Audiovisual → SEMPRE 'ver humano'" if is_av else "0 (não é rádio/TV/podcast)"
    c15_valor = '"ver humano"' if is_av else "0"

    titulo = str(row_data[0] or "")
    prompt = PROMPT_AVALIACAO.format(
        empresa=empresa, titulo=titulo, veiculo=veiculo, midia=midia,
        texto=texto, c15_instrucao=c15_instrucao, c15_valor=c15_valor,
    )

    try:
        resp = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=1000,
            messages=[{"role": "user", "content": prompt}],
        )
        txt = resp.content[0].text.strip()
        txt = re.sub(r"```json|```", "", txt).strip()
        result = json.loads(txt)

        extras = sum(
            result.get(f"c{n}", 0) if isinstance(result.get(f"c{n}"), (int, float)) else 0
            for n in [8, 9, 10, 11, 12, 13]
        )
        result["soma"] = (1 if result.get("protagonismo") == "Destaque" else 0) + extras
        return result

    except Exception as e:
        return {
            "protagonismo": "Menção", "crit_ativado": "nenhum",
            "c8": 0, "c9": 0, "c10": 0, "c11": 0, "c12": 0, "c13": 0,
            "c14": "ver humano",
            "c15": "ver humano" if is_av else 0,
            "c16": "ver humano", "soma": 0,
            "pr_tec_produto": False, "pr_puro_imobis": False,
            "data_proativa": False, "data_carona": False,
            "tema": "", "subtema": "", "produto": "",
            "obs": f"Erro na avaliação: {str(e)[:100]}. Avaliar manualmente.",
            "cidade": "", "confianca": "baixa",
        }


def gerar_whatsapp(por_empresa: dict) -> str:
    hoje = datetime.now().strftime("%d/%m")
    main = ["Loft", "QuintoAndar", "ZAP", "Superlógica"]
    emojis = {"Loft": "🟠", "QuintoAndar": "🔵", "ZAP": "🟢", "Superlógica": "⚫"}

    msg = f"📊 PR Score — {hoje}\n\n"
    for emp in main:
        d = por_empresa.get(emp, {"pts": 0, "mats": []})
        msg += f"{emojis.get(emp, '⚪')} {emp}: {d['pts']} pts\n"

    loft = por_empresa.get("Loft", {})
    if loft.get("mats"):
        msg += "\n📰 Destaques Loft do dia:\n"
        for m in loft["mats"]:
            msg += f"• {m['titulo']} - {m['veiculo']}\n{m['link']}\n"

    extras = [(e, d) for e, d in por_empresa.items() if e not in main and d["pts"] > 10]
    if extras:
        msg += "\n⚠️ Destaque secundário:\n"
        for emp, d in extras:
            msg += f"{emp}: {d['pts']} pts\n"

    return msg


def gerar_html_email(por_empresa: dict) -> str:
    hoje = datetime.now().strftime("%d/%m")
    main = ["Loft", "QuintoAndar", "ZAP", "Superlógica"]

    html = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head><meta charset="UTF-8">
<style>
body{{font-family:Arial,sans-serif;max-width:600px;margin:0 auto;padding:24px;color:#111;background:#fff}}
.hdr{{background:#F05A28;padding:20px 24px;border-radius:8px;margin-bottom:24px}}
.hdr h1{{color:#fff;font-size:20px;margin:0}}
.sec{{margin-bottom:24px}}
h2{{font-size:15px;font-weight:600;border-bottom:2px solid #F05A28;padding-bottom:6px;margin:0 0 10px}}
.item{{background:#FAFAF8;border-radius:6px;padding:10px 14px;margin:6px 0}}
.item a{{color:#F05A28;text-decoration:none;font-weight:500;font-size:14px}}
.item span{{display:block;color:#888;font-size:12px;margin-top:2px}}
</style>
</head>
<body>
<div class="hdr"><h1>📊 PR Score — {hoje}</h1></div>
"""
    for emp in main:
        d = por_empresa.get(emp, {})
        if not d or not d.get("mats"):
            continue
        html += f'<div class="sec"><h2>{emp} — {d["pts"]} pts</h2>\n'
        for m in d["mats"]:
            html += f'<div class="item"><a href="{m["link"]}">{m["titulo"]}</a><span>{m["veiculo"]}</span></div>\n'
        html += "</div>\n"

    html += "</body></html>"
    return html


def main():
    if len(sys.argv) < 2:
        print("Uso: python pr_score.py <caminho_para_planilha.xlsx>")
        sys.exit(1)

    arquivo = sys.argv[1]
    client = anthropic.Anthropic()  # lê ANTHROPIC_API_KEY do ambiente

    print(f"\n{'='*60}")
    print(f"PR Score — Processando: {arquivo}")
    print(f"{'='*60}\n")

    wb_orig = load_workbook(arquivo)

    # Tentar usar a aba "Matérias", se não existir usa a aba ativa
    if "Matérias" in wb_orig.sheetnames:
        ws_orig = wb_orig["Matérias"]
    else:
        ws_orig = wb_orig.active
        print(f"⚠  Aba 'Matérias' não encontrada — usando aba '{ws_orig.title}'")

    rows = list(ws_orig.iter_rows(values_only=True))
    header = list(rows[0])
    data = [list(r) for r in rows[1:] if any(c for c in r)]

    print(f"📋 {len(data)} linhas encontradas na planilha\n")

    # ── Classificar linhas ────────────────────────────────────────────────────
    valid_rows = []
    red_rows = []

    for i, row in enumerate(data):
        veiculo = str(row[1] or "").strip()
        if not veiculo:
            red_rows.append((i, "Veículo vazio"))
        elif veiculo.lower() in ("portas.com.br", "portas"):
            red_rows.append((i, "portas.com.br não é Tier 1"))
        elif veiculo not in TIER1:
            red_rows.append((i, f'Veículo "{veiculo}" não está na lista Tier 1'))
        else:
            valid_rows.append(i)

    print(f"🔴 {len(red_rows)} linhas não-Tier 1 (serão pintadas de vermelho)")
    print(f"✅ {len(valid_rows)} matérias Tier 1 para avaliar\n")

    # ── Avaliar matérias Tier 1 ───────────────────────────────────────────────
    gab_rows = []
    yellow_rows = []
    por_empresa = {}

    for idx_num, orig_idx in enumerate(valid_rows):
        row = data[orig_idx]
        titulo   = str(row[0] or "")
        veiculo  = str(row[1] or "").strip()
        canal    = str(row[6] or "").strip()
        estado   = str(row[11] or "")
        midia    = str(row[12] or "").strip()
        url_fonte = str(row[26] or "").strip()
        link     = str(row[27] or "").strip()
        data_str = excel_date(row[2])
        mes      = int(data_str.split("/")[1]) if data_str and "/" in data_str else datetime.now().month

        empresa = BRAND_MAP.get(canal.lower(), canal or veiculo)

        print(f"[{idx_num+1}/{len(valid_rows)}] {titulo[:70]}...")
        print(f"         Veículo: {veiculo} | Empresa: {empresa}")

        av = avaliar_materia(client, row, empresa, veiculo, midia, url_fonte, link)

        print(f"         → {av['protagonismo']} | Soma: {av['soma']} | Confiança: {av.get('confianca','?')}")
        if av.get("obs"):
            print(f"         ⚠  {av['obs'][:80]}")

        if av["protagonismo"] == "Destaque":
            if empresa not in por_empresa:
                por_empresa[empresa] = {"pts": 0, "mats": []}
            por_empresa[empresa]["pts"] += av["soma"]
            por_empresa[empresa]["mats"].append({
                "titulo": titulo, "veiculo": veiculo, "link": link
            })

        gab_row = [
            idx_num + 1, mes, titulo, veiculo, data_str, link, empresa,
            av["protagonismo"], av.get("crit_ativado", ""),
            av.get("c8", 0), av.get("c9", 0), av.get("c10", 0),
            av.get("c11", 0), av.get("c12", 0), av.get("c13", 0),
            av.get("c14", "ver humano"), av.get("c15", 0), av.get("c16", "ver humano"),
            av["soma"],
            str(av.get("pr_tec_produto", False)), str(av.get("pr_puro_imobis", False)),
            str(av.get("data_proativa", False)), str(av.get("data_carona", False)),
            "",  # Retranca — humano preenche
            av.get("tema", ""), av.get("subtema", ""), av.get("produto", ""),
            av.get("obs", ""), av.get("cidade", ""), estado,
        ]
        gab_rows.append(gab_row)

        if av.get("obs", "").strip():
            yellow_rows.append(len(gab_rows))  # índice 1-based na aba gabarito

        time.sleep(0.3)

    # ── Montar planilha de saída ───────────────────────────────────────────────
    print("\n📊 Montando planilha de saída...")

    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    # ── ABA MATÉRIAS ──────────────────────────────────────────────────────────
    ws_mat = wb_out.create_sheet("Matérias")
    mat_header = header + ["OBS — Motivo de exclusão"]
    ws_mat.append(mat_header)
    for cell in ws_mat[1]:
        cell.fill = ORANGE_FILL
        cell.font = HEADER_FONT

    red_set    = {i for i, _ in red_rows}
    red_motivos = {i: m for i, m in red_rows}

    for i, row in enumerate(data):
        row_com_obs = list(row) + [red_motivos.get(i, "")]
        ws_mat.append(row_com_obs)
        xlsx_row = i + 2
        if i in red_set:
            for cell in ws_mat[xlsx_row]:
                cell.fill = RED_FILL

    # ── ABA CÓPIA ─────────────────────────────────────────────────────────────
    ws_copy = wb_out.create_sheet("Cópia de Matérias")
    ws_copy.append(header)
    for cell in ws_copy[1]:
        cell.fill = ORANGE_FILL
        cell.font = HEADER_FONT
    for row in data:
        ws_copy.append(list(row))

    # ── ABA GABARITO DASHBOARD ────────────────────────────────────────────────
    ws_gab = wb_out.create_sheet("Gabarito Dashboard")
    ws_gab.append(GAB_HEADER)
    for cell in ws_gab[1]:
        cell.fill = ORANGE_FILL
        cell.font = HEADER_FONT

    for i, gab_row in enumerate(gab_rows):
        ws_gab.append(gab_row)
        if (i + 1) in yellow_rows:
            for cell in ws_gab[i + 2]:
                cell.fill = YELLOW_FILL

    # ── ABA TOTAIS ────────────────────────────────────────────────────────────
    ws_tot = wb_out.create_sheet("Totais")
    ws_tot.append(["PONTUAÇÃO POR EMPRESA"])
    ws_tot.append(["Empresa", "Matérias Tier 1", "Destaques", "Pontuação"])

    main_emps = ["Loft", "QuintoAndar", "ZAP", "Superlógica"]
    emp_stats = {}
    for gab_row in gab_rows:
        emp  = gab_row[6]
        prot = gab_row[7]
        soma = gab_row[18] if isinstance(gab_row[18], (int, float)) else 0
        if emp not in emp_stats:
            emp_stats[emp] = {"n": 0, "d": 0, "pts": 0}
        emp_stats[emp]["n"] += 1
        if prot == "Destaque":
            emp_stats[emp]["d"] += 1
            emp_stats[emp]["pts"] += soma

    for emp in [*main_emps, *[e for e in emp_stats if e not in main_emps]]:
        if emp in emp_stats:
            s = emp_stats[emp]
            ws_tot.append([emp, s["n"], s["d"], s["pts"]])

    ws_tot.append([])
    ws_tot.append(["POR ESTADO"])
    ws_tot.append(["Estado", "Matérias"])
    est_stats = {}
    for gab_row in gab_rows:
        est = gab_row[29] or "N/D"
        est_stats[est] = est_stats.get(est, 0) + 1
    for k, v in sorted(est_stats.items(), key=lambda x: -x[1]):
        ws_tot.append([k, v])

    # ── Salvar ────────────────────────────────────────────────────────────────
    data_hoje = datetime.now().strftime("%Y-%m-%d")
    output_path = f"PR_Score_{data_hoje}.xlsx"
    wb_out.save(output_path)
    print(f"\n✅ Planilha salva: {output_path}")

    wpp  = gerar_whatsapp(por_empresa)
    html = gerar_html_email(por_empresa)

    wpp_path  = f"PR_Score_{data_hoje}_whatsapp.txt"
    html_path = f"PR_Score_{data_hoje}_email.html"

    with open(wpp_path, "w", encoding="utf-8") as f:
        f.write(wpp)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)

    print(f"✅ WhatsApp salvo: {wpp_path}")
    print(f"✅ HTML do e-mail salvo: {html_path}")

    # ── Resumo final ──────────────────────────────────────────────────────────
    print(f"\n{'='*60}")
    print(f"RESUMO DO DIA")
    print(f"{'='*60}")
    print(f"📋 Matérias na planilha:    {len(data)}")
    print(f"🔴 Pintadas de vermelho:    {len(red_rows)}")
    print(f"✅ Avaliadas (Tier 1):      {len(valid_rows)}")
    print(f"⚠️  Precisam olhar humano:  {len(yellow_rows)}")
    print(f"\n📊 PONTUAÇÃO DO DIA:")
    emoji_map = {"Loft": "🟠", "QuintoAndar": "🔵", "ZAP": "🟢", "Superlógica": "⚫"}
    for emp in main_emps:
        pts  = por_empresa.get(emp, {}).get("pts", 0)
        mats = por_empresa.get(emp, {}).get("mats", [])
        emoji = emoji_map.get(emp, "⚪")
        print(f"   {emoji} {emp}: {pts} pts ({len(mats)} destaques)")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    main()
